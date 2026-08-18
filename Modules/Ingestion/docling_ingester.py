from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

import tiktoken
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter

_PAGE_MARKER_PATTERN = re.compile(r"<!-- PAGE (\d+) -->", re.IGNORECASE)
_IMAGE_MARKER_PATTERN = re.compile(r"(<!-- image -->|!\[Image\]\([^)]+\))", re.IGNORECASE)
_TOKEN_PATTERN = re.compile(r"\w+|[^\w\s]", re.UNICODE)


class _FallbackTokenEncoder:
    @staticmethod
    def encode(text: str) -> List[str]:
        return _TOKEN_PATTERN.findall(text)


@dataclass(slots=True)
class PdfPipelineConfig:
    do_ocr: bool = True
    force_full_page_ocr: bool = True
    do_table_structure: bool = True
    images_scale: float = 1.0
    layout_batch_size: int = 2
    ocr_batch_size: int = 4
    table_batch_size: int = 2
    allow_external_plugins: bool = False


@dataclass(slots=True)
class ChunkingConfig:
    chunk_size_tokens: int = 1024
    chunk_overlap_tokens: int = 256
    preview_characters: int = 240
    separators: Sequence[str] = field(default_factory=lambda: ("\n\n", "\n", " ", ""))
    headers_to_split_on: Sequence[tuple[str, str]] = field(
        default_factory=lambda: (
            ("#", "H1"),
            ("##", "H2"),
            ("###", "H3"),
            ("####", "H4"),
            ("#####", "H5"),
            ("######", "H6"),
        )
    )


@dataclass(slots=True)
class ParsedDocumentChunk:
    chunk_id: str
    content: str
    metadata: Dict[str, Any]


@dataclass(slots=True)
class ParsedDocument:
    document_id: str
    source: str
    file_name: str
    file_path: str
    parser_name: str
    content: str
    page_count: Optional[int]
    chunks: List[ParsedDocumentChunk]
    strategy_name: str = "auto"
    selected_mode: str = "auto"
    selection_reason: str = ""
    fallback_used: bool = False
    analysis: Dict[str, Any] = field(default_factory=dict)
    capabilities: Dict[str, Any] = field(default_factory=dict)
    warnings: List[str] = field(default_factory=list)
    attempted_modes: List[str] = field(default_factory=list)
    mode_errors: List[str] = field(default_factory=list)


class DocumentIngestor:
    def __init__(
        self,
        pdf_config: Optional[PdfPipelineConfig] = None,
        chunking_config: Optional[ChunkingConfig] = None,
        encoding_name: str = "cl100k_base",
    ) -> None:
        self.pdf_config = pdf_config or PdfPipelineConfig()
        self.chunking_config = chunking_config or ChunkingConfig()
        self._token_encoder = self._build_token_encoder(encoding_name)
        self._pdf_converter = None

    def parse_file(
        self,
        file_path: str | Path,
        *,
        source: Optional[str] = None,
        document_id: Optional[str] = None,
    ) -> ParsedDocument:
        return self._parse_path(
            file_path=file_path,
            source=source,
            document_id=document_id,
            pdf_mode="auto",
        )

    def parse_file_as_text(
        self,
        file_path: str | Path,
        *,
        source: Optional[str] = None,
        document_id: Optional[str] = None,
    ) -> ParsedDocument:
        return self._parse_path(
            file_path=file_path,
            source=source,
            document_id=document_id,
            pdf_mode="text_only",
        )

    def parse_file_with_docling(
        self,
        file_path: str | Path,
        *,
        source: Optional[str] = None,
        document_id: Optional[str] = None,
    ) -> ParsedDocument:
        return self._parse_path(
            file_path=file_path,
            source=source,
            document_id=document_id,
            pdf_mode="docling_only",
        )

    def _parse_path(
        self,
        *,
        file_path: str | Path,
        source: Optional[str],
        document_id: Optional[str],
        pdf_mode: str,
    ) -> ParsedDocument:
        path = Path(file_path).expanduser()
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")
        if not path.is_file():
            raise ValueError(f"Path is not a file: {path}")

        content, parser_name, page_count, use_markdown_headers = self._read_document(
            path,
            pdf_mode=pdf_mode,
        )
        cleaned_content = content.strip()
        if not cleaned_content:
            raise ValueError(f"No readable text was found in {path.name}.")

        resolved_document_id = document_id or path.stem
        resolved_source = source or path.name
        chunks = self._build_chunks(
            content=cleaned_content,
            document_id=resolved_document_id,
            source=resolved_source,
            file_name=path.name,
            file_path=str(path),
            document_name=path.stem,
            parser_name=parser_name,
            page_count=page_count,
            use_markdown_headers=use_markdown_headers,
        )

        return ParsedDocument(
            document_id=resolved_document_id,
            source=resolved_source,
            file_name=path.name,
            file_path=str(path),
            parser_name=parser_name,
            content=cleaned_content,
            page_count=page_count,
            chunks=chunks,
        )

    def _read_document(
        self,
        path: Path,
        *,
        pdf_mode: str = "auto",
    ) -> tuple[str, str, Optional[int], bool]:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            if pdf_mode == "text_only":
                return self._read_pdf(path), "pdf_text", self._count_pdf_pages(path), False
            if pdf_mode == "docling_only":
                markdown, page_count = self._convert_pdf_to_markdown(path)
                return markdown, "docling_pdf", page_count, True
            try:
                markdown, page_count = self._convert_pdf_to_markdown(path)
                return markdown, "docling_pdf", page_count, True
            except Exception:
                return self._read_pdf(path), "pypdf_fallback", self._count_pdf_pages(path), False
        if suffix == ".docx":
            return self._read_docx(path), "docx_text", None, False
        if suffix in {".md", ".markdown"}:
            return self._read_text_file(path), "markdown_text", None, True
        return self._read_text_file(path), "text", None, False

    def _build_chunks(
        self,
        *,
        content: str,
        document_id: str,
        source: str,
        file_name: str,
        file_path: str,
        document_name: str,
        parser_name: str,
        page_count: Optional[int],
        use_markdown_headers: bool,
    ) -> List[ParsedDocumentChunk]:
        splitter = self._build_recursive_splitter()
        raw_chunks = []
        if use_markdown_headers:
            header_splitter = MarkdownHeaderTextSplitter(
                headers_to_split_on=list(self.chunking_config.headers_to_split_on),
                strip_headers=False,
            )
            for section in self._split_markdown_sections(content):
                normalized = self._normalize_markdown(section["content"])
                header_documents = header_splitter.split_text(normalized)
                for raw_chunk in splitter.split_documents(header_documents):
                    raw_chunk.metadata.setdefault(
                        "_default_page_start",
                        section["page_start"],
                    )
                    raw_chunk.metadata.setdefault(
                        "_default_page_end",
                        section["page_end"],
                    )
                    raw_chunks.append(raw_chunk)
        else:
            raw_chunks = splitter.create_documents([content])

        prepared_chunks = []
        for raw_chunk in raw_chunks:
            page_start, page_end, cleaned_chunk = self._extract_page_span(raw_chunk.page_content)
            default_page_start = raw_chunk.metadata.pop("_default_page_start", None)
            default_page_end = raw_chunk.metadata.pop("_default_page_end", None)
            if page_start is None and cleaned_chunk:
                page_start = default_page_start
                page_end = default_page_end
            if not cleaned_chunk:
                continue
            prepared_chunks.append(
                {
                    "content": cleaned_chunk,
                    "headers": self._join_headers(raw_chunk.metadata),
                    "page_start": page_start,
                    "page_end": page_end,
                    "token_count": self._count_tokens(cleaned_chunk),
                }
            )

        if not prepared_chunks:
            page_start, page_end, cleaned_chunk = self._extract_page_span(content)
            prepared_chunks.append(
                {
                    "content": cleaned_chunk,
                    "headers": "",
                    "page_start": page_start,
                    "page_end": page_end,
                    "token_count": self._count_tokens(cleaned_chunk),
                }
            )

        total_chunks = len(prepared_chunks)
        chunks: List[ParsedDocumentChunk] = []
        for chunk_index, prepared_chunk in enumerate(prepared_chunks):
            chunk_id = f"{document_id}:chunk:{chunk_index:04d}"
            metadata = self._build_chunk_metadata(
                source=source,
                file_name=file_name,
                file_path=file_path,
                document_id=document_id,
                document_name=document_name,
                parser_name=parser_name,
                page_count=page_count,
                chunk_id=chunk_id,
                chunk_index=chunk_index,
                chunk_count=total_chunks,
                headers=prepared_chunk["headers"],
                page_start=prepared_chunk["page_start"],
                page_end=prepared_chunk["page_end"],
                token_count=prepared_chunk["token_count"],
                content=prepared_chunk["content"],
            )
            chunks.append(
                ParsedDocumentChunk(
                    chunk_id=chunk_id,
                    content=prepared_chunk["content"],
                    metadata=metadata,
                )
            )

        return chunks

    def _build_chunk_metadata(
        self,
        *,
        source: str,
        file_name: str,
        file_path: str,
        document_id: str,
        document_name: str,
        parser_name: str,
        page_count: Optional[int],
        chunk_id: str,
        chunk_index: int,
        chunk_count: int,
        headers: str,
        page_start: Optional[int],
        page_end: Optional[int],
        token_count: int,
        content: str,
    ) -> Dict[str, Any]:
        metadata = {
            "source": source,
            "ingest_type": "file_chunk",
            "parser": parser_name,
            "fileName": file_name,
            "filePath": file_path,
            "documentId": document_id,
            "documentName": document_name,
            "docName": document_name,
            "chunkId": chunk_id,
            "chunkNum": chunk_index,
            "chunkCount": chunk_count,
            "tokenCount": token_count,
            "context": self._build_preview(content),
            "headers": headers,
            "pageStart": page_start,
            "pageEnd": page_end,
            "pageCount": page_count,
        }
        return {
            key: value
            for key, value in metadata.items()
            if value is not None and value != ""
        }

    def _convert_pdf_to_markdown(self, path: Path) -> tuple[str, Optional[int]]:
        converter = self._get_pdf_converter()
        results = list(converter.convert_all([str(path)]))
        if not results:
            raise RuntimeError(f"Docling did not return a conversion result for {path.name}.")

        result = results[0]

        from docling_core.types.doc import ImageRefMode

        page_map = getattr(result.document, "pages", {}) or {}
        page_numbers = sorted(int(page_number) for page_number in page_map.keys())
        if not page_numbers:
            markdown = result.document.export_to_markdown(
                image_mode=ImageRefMode.PLACEHOLDER,
                image_placeholder="<!-- image -->",
            )
            return markdown, None

        markdown_sections = []
        for page_number in page_numbers:
            page_markdown = result.document.export_to_markdown(
                page_no=page_number,
                image_mode=ImageRefMode.PLACEHOLDER,
                image_placeholder="<!-- image -->",
            ).strip()
            if page_markdown:
                markdown_sections.append(f"<!-- PAGE {page_number} -->\n{page_markdown}")

        markdown = "\n\n".join(markdown_sections).strip()
        return markdown, len(page_numbers)

    def _get_pdf_converter(self):
        if self._pdf_converter is not None:
            return self._pdf_converter

        try:
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.pipeline_options import (
                EasyOcrOptions,
                TableStructureOptions,
                ThreadedPdfPipelineOptions,
            )
            from docling.document_converter import DocumentConverter, PdfFormatOption
        except ImportError as exc:
            raise RuntimeError(
                "PDF ingestion requires Docling and its PDF dependencies in the environment."
            ) from exc

        pipeline_options = ThreadedPdfPipelineOptions()
        pipeline_options.images_scale = self.pdf_config.images_scale
        pipeline_options.do_table_structure = self.pdf_config.do_table_structure
        pipeline_options.table_structure_options = TableStructureOptions(
            do_cell_matching=True
        )
        pipeline_options.do_ocr = self.pdf_config.do_ocr
        pipeline_options.layout_batch_size = self.pdf_config.layout_batch_size
        pipeline_options.ocr_batch_size = self.pdf_config.ocr_batch_size
        pipeline_options.table_batch_size = self.pdf_config.table_batch_size
        pipeline_options.allow_external_plugins = self.pdf_config.allow_external_plugins
        if self.pdf_config.do_ocr:
            pipeline_options.ocr_options = EasyOcrOptions(
                force_full_page_ocr=self.pdf_config.force_full_page_ocr
            )

        self._pdf_converter = DocumentConverter(
            format_options={
                InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options)
            }
        )
        return self._pdf_converter

    def _build_recursive_splitter(self) -> RecursiveCharacterTextSplitter:
        return RecursiveCharacterTextSplitter(
            chunk_size=self.chunking_config.chunk_size_tokens,
            chunk_overlap=self.chunking_config.chunk_overlap_tokens,
            length_function=self._count_tokens,
            separators=list(self.chunking_config.separators),
        )

    def _count_tokens(self, text: str) -> int:
        return len(self._token_encoder.encode(text))

    def _build_preview(self, text: str) -> str:
        normalized = " ".join(text.split())
        if len(normalized) <= self.chunking_config.preview_characters:
            return normalized
        return normalized[: self.chunking_config.preview_characters].rstrip() + "..."

    @staticmethod
    def _join_headers(metadata: Dict[str, Any]) -> str:
        return " > ".join(str(value) for value in metadata.values() if value)

    @staticmethod
    def _normalize_markdown(text: str) -> str:
        normalized = _IMAGE_MARKER_PATTERN.sub(r"\n\n## [Figure]\n\1\n\n", text)
        normalized = _PAGE_MARKER_PATTERN.sub(
            lambda match: f"\n\n<!-- PAGE {match.group(1)} -->\n\n",
            normalized,
        )
        normalized = re.sub(r"\n{3,}", "\n\n", normalized)
        return normalized.strip()

    @staticmethod
    def _extract_page_span(text: str) -> tuple[Optional[int], Optional[int], str]:
        page_numbers = DocumentIngestor._find_page_numbers(text)
        cleaned = _PAGE_MARKER_PATTERN.sub("", text)
        cleaned = re.sub(r"<!-- image -->", "[Image]", cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
        if not page_numbers:
            return None, None, cleaned
        return min(page_numbers), max(page_numbers), cleaned

    @staticmethod
    def _find_page_numbers(text: str) -> List[int]:
        return [int(match.group(1)) for match in _PAGE_MARKER_PATTERN.finditer(text)]

    @staticmethod
    def _split_markdown_sections(text: str) -> List[Dict[str, Any]]:
        page_markers = list(_PAGE_MARKER_PATTERN.finditer(text))
        if not page_markers:
            return [{"page_start": None, "page_end": None, "content": text}]

        sections = []
        for index, marker in enumerate(page_markers):
            next_start = (
                page_markers[index + 1].start()
                if index + 1 < len(page_markers)
                else len(text)
            )
            section_text = text[marker.end() : next_start].strip()
            if not section_text:
                continue
            page_number = int(marker.group(1))
            sections.append(
                {
                    "page_start": page_number,
                    "page_end": page_number,
                    "content": section_text,
                }
            )

        if sections:
            return sections
        return [{"page_start": None, "page_end": None, "content": text}]

    @staticmethod
    def _read_text_file(path: Path) -> str:
        encodings = ("utf-8", "utf-8-sig", "cp1252", "latin-1")
        for encoding in encodings:
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Unable to decode {path.name} as text.")

    @staticmethod
    def _read_docx(path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "DOCX support requires the 'python-docx' package in the environment."
            ) from exc

        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        return "\n".join(line for line in paragraphs if line.strip())

    @staticmethod
    def _read_pdf(path: Path) -> str:
        try:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(path))
            extracted_pages = []
            for page_index in range(len(pdf)):
                text_page = pdf[page_index].get_textpage()
                extracted_pages.append(text_page.get_text_range() or "")
            return "\n\n".join(page for page in extracted_pages if page.strip())
        except ImportError:
            pass

        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError(
                "PDF fallback support requires either 'pypdfium2' or 'pypdf' in the environment."
            ) from exc

        reader = PdfReader(str(path))
        extracted_pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(page for page in extracted_pages if page.strip())

    @staticmethod
    def _count_pdf_pages(path: Path) -> Optional[int]:
        try:
            import pypdfium2 as pdfium

            return len(pdfium.PdfDocument(str(path)))
        except Exception:
            pass

        try:
            from pypdf import PdfReader

            return len(PdfReader(str(path)).pages)
        except Exception:
            return None

    @staticmethod
    def _build_token_encoder(encoding_name: str):
        try:
            return tiktoken.get_encoding(encoding_name)
        except Exception:
            try:
                return tiktoken.encoding_for_model(encoding_name)
            except Exception:
                pass
        try:
            return tiktoken.get_encoding("cl100k_base")
        except Exception:
            return _FallbackTokenEncoder()
